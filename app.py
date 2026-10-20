# -*- coding: utf-8 -*-
"""ChurnLens - Team Visionaries | run: streamlit run app.py"""
import base64
import io
import os
import numpy as np
import pandas as pd
import streamlit as st
from PIL import Image

import decision as D
from common import MULTI
from i18n import LANG_NAMES, T
from i18n_extra import XT

HERE = os.path.dirname(os.path.abspath(__file__))
LOGO = os.path.join(HERE, "logo.jpg")
PAGES = ["home", "analyze", "batch", "model", "about"]

st.set_page_config(page_title="ChurnLens | Team Visionaries",
                   page_icon=Image.open(LOGO) if os.path.exists(LOGO) else None, layout="wide")


@st.cache_data(show_spinner=False)
def logo_b64(width=420):
    """Small base64 copy of the logo for inline HTML."""
    if not os.path.exists(LOGO):
        return ""
    im = Image.open(LOGO).convert("RGB")
    im.thumbnail((width, width))
    buf = io.BytesIO()
    im.save(buf, "JPEG", quality=88)
    return base64.b64encode(buf.getvalue()).decode()


CSS = """
<style>
@import url('https://fonts.googleapis.com/css2?family=Poppins:wght@400;500;600;700;800&family=Noto+Sans+Devanagari:wght@400;600;700&display=swap');
.stApp{font-family:'Poppins','Noto Sans Devanagari',sans-serif;
  background:radial-gradient(900px 500px at 85% -10%,rgba(30,144,255,.16),transparent 60%),
             radial-gradient(700px 420px at -5% 10%,rgba(34,184,255,.10),transparent 60%),#000}
.stApp h1,.stApp h2,.stApp h3,.stApp h4,.stApp button,.stApp label{font-family:'Poppins','Noto Sans Devanagari',sans-serif}
header[data-testid="stHeader"]{background:transparent}
footer{visibility:hidden}
.block-container{max-width:1180px;padding-top:1.2rem;padding-bottom:2rem}

/* ---------- top bar ---------- */
.st-key-topbar{background:rgba(11,18,32,.88);border:1px solid #17304f;border-radius:16px;
  padding:.55rem 1.1rem;margin-bottom:1.4rem;box-shadow:0 10px 30px rgba(0,0,0,.45)}
.brand{display:flex;align-items:center;gap:.7rem}
.brand img{width:42px;height:42px;border-radius:10px;object-fit:cover}
.brand b{font-size:1.15rem;color:#fff;letter-spacing:.3px;display:block;line-height:1.15}
.brand small{color:#7f93ad;font-size:.72rem;display:block}
.st-key-lang,.st-key-page{display:flex;justify-content:center}
.st-key-lang{justify-content:flex-end}

/* ---------- buttons ---------- */
.stButton button[kind="primary"],.stButton button[data-testid="stBaseButton-primary"]{
  background:linear-gradient(90deg,#1e90ff,#22b8ff);border:0;border-radius:999px;font-weight:700;
  padding:.7rem 1.4rem;box-shadow:0 10px 26px rgba(30,144,255,.38);transition:transform .15s}
.stButton button[kind="primary"]:hover,.stButton button[data-testid="stBaseButton-primary"]:hover{transform:translateY(-2px)}

/* ---------- hero ---------- */
.hero2{position:relative;text-align:center;padding:1.2rem 0 .4rem}
.badge{display:inline-block;border:1px solid #1e4a7a;background:rgba(30,144,255,.12);color:#8fd0ff;
  border-radius:999px;padding:5px 16px;font-size:.8rem;font-weight:600;letter-spacing:.3px}
.hlogo{display:block;margin:1rem auto .2rem;width:190px;max-width:55%;border-radius:22px;
  box-shadow:0 0 70px rgba(30,144,255,.35)}
.hero2 h1{font-size:4rem;font-weight:800;margin:.3rem 0 0;color:#fff;letter-spacing:-.5px;line-height:1.1}
.hero2 h1 span{background:linear-gradient(90deg,#1e90ff,#22d3ee);-webkit-background-clip:text;
  background-clip:text;-webkit-text-fill-color:transparent}
.hero2 .tag{color:#22b8ff;font-size:1.2rem;font-weight:600;margin:.5rem auto 0;max-width:860px}
.hero2 .sub{color:#9fb0c6;font-size:1.02rem;max-width:720px;margin:.6rem auto 0}
.chips{text-align:center;margin:1rem 0 .4rem}
.chip{display:inline-block;border:1px solid #1e90ff;color:#22b8ff;border-radius:999px;
  padding:4px 15px;margin:3px;font-size:.84rem;font-weight:500}

/* ---------- section headings ---------- */
.sec{text-align:center;margin:3rem 0 1.3rem}
.sec h2{font-size:1.85rem;font-weight:700;color:#fff;margin:0}
.sec p{color:#8da0b8;margin:.35rem 0 0;font-size:.95rem}
.sec:after{content:"";display:block;width:56px;height:3px;border-radius:3px;margin:.8rem auto 0;
  background:linear-gradient(90deg,#1e90ff,#22d3ee)}

/* ---------- demo card ---------- */
.demo{max-width:760px;margin:0 auto;background:linear-gradient(160deg,#0d1a2e,#08101d);
  border:1px solid #1b3a60;border-radius:22px;padding:1.5rem 1.7rem;box-shadow:0 20px 60px rgba(0,0,0,.55)}
.demo-top{display:flex;align-items:center;gap:1.4rem;flex-wrap:wrap}
.ring{width:104px;height:104px;border-radius:50%;flex:none;display:flex;align-items:center;justify-content:center;
  background:conic-gradient(#d64545 0 78%,#1c2a3f 78% 100%)}
.ring div{width:80px;height:80px;border-radius:50%;background:#0a1322;display:flex;align-items:center;
  justify-content:center;font-size:1.65rem;font-weight:800;color:#fff}
.demo-top .meta{flex:1;min-width:200px;text-align:left}
.demo-top .meta small{color:#8da0b8;font-size:.82rem}
.demo-top .meta .pill{display:inline-block;margin-top:.35rem;padding:5px 16px;border-radius:999px;
  background:#d64545;color:#fff;font-weight:700;font-size:.82rem}
.trust{border:1px solid #1e4a7a;border-radius:14px;padding:.55rem 1rem;text-align:center;background:rgba(30,144,255,.08)}
.trust b{display:block;font-size:1.35rem;color:#22b8ff}
.trust small{color:#8da0b8;font-size:.72rem}
.bars{margin-top:1.2rem;text-align:left}
.barrow{display:flex;align-items:center;gap:.8rem;margin:.45rem 0;font-size:.85rem;color:#c3cfdf}
.barrow .nm{width:44%;white-space:nowrap;overflow:hidden;text-overflow:ellipsis}
.barrow .tr{flex:1;height:10px;background:#16233a;border-radius:6px;overflow:hidden}
.barrow .fl{height:100%;border-radius:6px}
.up{background:linear-gradient(90deg,#d64545,#ff7a7a)}.dn{background:linear-gradient(90deg,#1f9d55,#4be08a)}
.wi{margin-top:1rem;background:rgba(31,157,85,.10);border:1px solid rgba(31,157,85,.45);color:#9be7bd;
  border-radius:12px;padding:.6rem 1rem;font-size:.88rem;text-align:left}

/* ---------- steps ---------- */
.steps{display:grid;grid-template-columns:repeat(3,1fr);gap:1.1rem;max-width:960px;margin:0 auto}
.step{text-align:center;padding:1rem .8rem}
.step .n{width:46px;height:46px;border-radius:50%;margin:0 auto .7rem;display:flex;align-items:center;
  justify-content:center;font-weight:800;color:#fff;background:linear-gradient(135deg,#1e90ff,#22d3ee);
  box-shadow:0 8px 22px rgba(30,144,255,.4)}
.step h4{margin:0 0 .25rem;color:#fff;font-size:1.05rem}
.step p{margin:0;color:#9fb0c6;font-size:.88rem}

/* ---------- feature cards ---------- */
.grid{display:grid;grid-template-columns:repeat(3,1fr);gap:1.1rem}
.card{background:linear-gradient(160deg,#0d1a2e,#09111f);border:1px solid #17304f;border-radius:18px;
  padding:1.3rem 1.3rem 1.2rem;transition:transform .2s,border-color .2s,box-shadow .2s}
.card:hover{transform:translateY(-4px);border-color:#1e90ff;box-shadow:0 14px 34px rgba(30,144,255,.18)}
.card .ic{width:46px;height:46px;border-radius:13px;display:flex;align-items:center;justify-content:center;
  font-size:1.35rem;background:rgba(30,144,255,.14);border:1px solid #1e4a7a;margin-bottom:.8rem}
.card h4{margin:0 0 .35rem;color:#22b8ff;font-size:1.05rem;font-weight:700}
.card p{margin:0;color:#c3cfdf;font-size:.9rem;line-height:1.5}
@media(max-width:820px){.grid,.steps{grid-template-columns:1fr}.hero2 h1{font-size:2.8rem}}

/* ---------- misc (other pages) ---------- */
.pill{display:inline-block;padding:6px 16px;border-radius:999px;font-weight:700;color:#fff;margin:4px 0}
.low{background:#1f9d55}.medium{background:#d98e04}.high{background:#d64545}
.note{background:#0b1220;border:1px solid #17304f;border-radius:12px;padding:12px 16px;color:#9fb0c6;
  margin-top:1rem;text-align:center}
.foot{text-align:center;color:#6f819a;font-size:.8rem;margin-top:2.5rem;border-top:1px solid #17304f;padding-top:.9rem}

/* ---------- decision card + kpi ---------- */
.dec{border-radius:18px;padding:1.1rem 1.4rem;margin:.2rem 0 1.1rem;border:1px solid;box-shadow:0 14px 40px rgba(0,0,0,.5)}
.dec-t{font-size:.78rem;letter-spacing:.6px;text-transform:uppercase;opacity:.8}
.dec-n{font-size:1.7rem;font-weight:800;margin:.1rem 0 .25rem;line-height:1.2}
.dec-b{font-size:.95rem;opacity:.95}
.dec-act{background:linear-gradient(135deg,#0c2a1a,#0a1f14);border-color:#1f9d55;color:#9be7bd}
.dec-review{background:linear-gradient(135deg,#2b2108,#1d1706);border-color:#d98e04;color:#ffd27a}
.dec-distrust{background:linear-gradient(135deg,#2c1010,#1d0a0a);border-color:#d64545;color:#ffaaaa}
.dec-monitor{background:linear-gradient(135deg,#0b1d33,#08121f);border-color:#1e90ff;color:#9fd0ff}
.dec-no_roi{background:linear-gradient(135deg,#1a2230,#111826);border-color:#6b7a90;color:#c3cfdf}
.newb{background:linear-gradient(90deg,#1f9d55,#22b86b);border-color:#1f9d55;color:#fff}
.mrow{display:grid;grid-template-columns:repeat(4,1fr);gap:.6rem;margin-top:1rem}
.mk{background:#0b1220;border:1px solid #17304f;border-radius:12px;padding:.5rem .7rem;text-align:left}
.mk small{display:block;color:#8da0b8;font-size:.7rem}.mk b{color:#fff;font-size:1.15rem}.mk.good b{color:#4be08a}
.kpi{background:#0b1220;border:1px solid #17304f;border-radius:14px;padding:.7rem 1rem;margin-bottom:.4rem}
.kpi small{display:block;color:#8da0b8;font-size:.76rem}
.kpi b{font-size:1.6rem;color:#fff}
</style>
"""
st.markdown(CSS, unsafe_allow_html=True)

# ------------------------------------------------------------------ language + nav (top bar)
if st.session_state.get("lang") not in LANG_NAMES:
    st.session_state.lang = "en"
if st.session_state.get("page") not in PAGES:
    st.session_state.page = "home"

with st.container(key="topbar"):
    c_brand, c_nav, c_lang = st.columns([2.3, 4.6, 2.4], vertical_alignment="center")
    with c_lang:
        st.segmented_control("Language", list(LANG_NAMES), key="lang", label_visibility="collapsed",
                             format_func=lambda c: LANG_NAMES[c])
    lang = st.session_state.get("lang") or "en"
    L = {**T[lang], **XT[lang]}
    L["nav"] = {**T[lang]["nav"], **XT[lang]["nav"]}
    with c_nav:
        st.segmented_control("Menu", PAGES, key="page", label_visibility="collapsed",
                             format_func=lambda c: L["nav"][c])
    with c_brand:
        st.markdown("<div class='brand'><img src='data:image/jpeg;base64,%s'><div><b>ChurnLens</b>"
                    "<small>Team Visionaries</small></div></div>" % logo_b64(120), unsafe_allow_html=True)

PAGE = st.session_state.get("page") or "home"


def vl(v):
    return L["vals"].get(v, v)


def flabel(f):
    if "_" in f:
        base, val = f.split("_", 1)
        if base in MULTI:
            return "%s: %s" % (L["feat"].get(base, base), vl(val))
    return L["feat"].get(f, f)


@st.cache_resource(show_spinner=False)
def get_bundle():
    from engine import load_bundle
    return load_bundle()


@st.cache_resource(show_spinner=False)
def get_explainers():
    from engine import Explainers
    return Explainers(get_bundle())


def go_analyze():
    st.session_state.page = "analyze"


def go_batch():
    st.session_state.page = "batch"


# ------------------------------------------------------------------ HOME
ICONS = ["🔍", "🛡️", "🚦", "💰", "🎛️", "📂", "🗺️", "🗣️", "⚖️"]


def page_home():
    g = lambda k, d: L.get(k, d)  # noqa: E731  (new keys fall back to English)
    hero = (
        "<div class='hero2'>"
        "<span class='badge'>%s</span> <span class='badge newb'>%s</span>"
        "<img class='hlogo' src='data:image/jpeg;base64,%s'>"
        "<h1>Churn<span>Lens</span></h1>"
        "<div class='tag'>%s</div><div class='sub'>%s</div></div>"
        % (g("badge", T["en"]["badge"]), L["hm_new"], logo_b64(420), L["tagline"], L["hero"]))
    st.markdown(hero, unsafe_allow_html=True)

    b1, b2, b3, b4 = st.columns([1.4, 1.3, 1.3, 1.4])
    with b2:
        st.button(L["start"], type="primary", width="stretch", on_click=go_analyze)
    with b3:
        st.button("📂 " + L["nav"]["batch"], width="stretch", on_click=go_batch)
    st.markdown("<div class='chips'>%s</div>" % "".join("<span class='chip'>%s</span>" % s for s in L["stats"]),
                unsafe_allow_html=True)

    # --- sample result card (static, illustrative numbers)
    F = L["feat"]
    rows = [
        ("%s: %s" % (F["Contract"], vl("Month-to-month")), 92, "up"),
        ("%s" % F["tenure"], 74, "up"),
        ("%s: %s" % (F["PaymentMethod"], vl("Electronic check")), 48, "up"),
        ("%s" % F["Partner"], 26, "dn"),
    ]
    bars = "".join("<div class='barrow'><div class='nm'>%s</div><div class='tr'><div class='fl %s' "
                   "style='width:%d%%'></div></div></div>" % (n, c, w) for n, w, c in rows)
    kp = lambda lab, val, cls="": "<div class='mk %s'><small>%s</small><b>%s</b></div>" % (cls, lab, val)  # noqa: E731
    wc = L["wi_cols2"]
    money_row = ("<div class='mrow'>%s%s%s%s</div>" % (
        kp(L["m_rar"], "$796"), kp(wc[3], "$102"), kp(wc[4], "$265"), kp(L["lv_net"], "+$163", "good")))
    demo = (
        "<div class='sec'><h2>%s</h2><p>%s</p></div>"
        "<div class='demo'><div class='dec dec-act' style='margin:0 0 1rem'><div class='dec-t'>%s</div>"
        "<div class='dec-n'>%s</div><div class='dec-b'>%s</div></div>"
        "<div class='demo-top'>"
        "<div class='ring'><div>78%%</div></div>"
        "<div class='meta'><small>%s</small><br><span class='pill'>%s</span></div>"
        "<div class='trust'><b>86/100</b><small>%s</small></div></div>"
        "%s<div class='bars'>%s</div>"
        "<div class='wi'>💡 %s: <b>78%% → 52%%</b></div></div>"
        % (g("demo_title", T["en"]["demo_title"]), g("demo_sub", T["en"]["demo_sub"]),
           L["dec_title"], L["dec_names"]["act"], L["dec_body"]["act"],
           L["prob"], L["risk"]["high"], L["tr_title"], money_row, bars, L["scen"]["contract_1y"]))
    st.markdown(demo, unsafe_allow_html=True)

    # --- how it works
    steps = g("steps", T["en"]["steps"])
    st.markdown(
        "<div class='sec'><h2>%s</h2></div><div class='steps'>%s</div>"
        % (g("how_title", T["en"]["how_title"]),
           "".join("<div class='step'><div class='n'>%d</div><h4>%s</h4><p>%s</p></div>" % (i + 1, t, d)
                   for i, (t, d) in enumerate(steps))), unsafe_allow_html=True)

    # --- feature cards
    cards = "".join("<div class='card'><div class='ic'>%s</div><h4>%s</h4><p>%s</p></div>" % (ICONS[i], t, d)
                    for i, (t, d) in enumerate(L["feats"]))
    st.markdown("<div class='sec'><h2>%s</h2></div><div class='grid'>%s</div>" % (L["feats_title"], cards),
                unsafe_allow_html=True)

    st.markdown("<div class='note'>🔒 %s</div>" % L["note"], unsafe_allow_html=True)


# ------------------------------------------------------------------ helpers for the decision tool
FEAT_ICON = {
    "Contract": "📄", "tenure": "⏳", "PaymentMethod": "💳", "TechSupport": "🛠️", "OnlineSecurity": "🔒",
    "InternetService": "🌐", "MonthlyCharges": "💰", "TotalCharges": "🧾", "Partner": "👫", "Dependents": "👪",
    "SeniorCitizen": "🧓", "PaperlessBilling": "📧", "PhoneService": "📞", "MultipleLines": "☎️",
    "StreamingTV": "📺", "StreamingMovies": "🎬", "OnlineBackup": "☁️", "DeviceProtection": "🛡️",
}
DEC_ORDER = ["act", "review", "distrust", "no_roi", "monitor"]
VOICE_LANG = {"en": "en-US", "hn": "en-IN", "hi": "hi-IN"}
DEFAULT_NAME = {"en": "Customer", "hn": "Customer", "hi": "ग्राहक"}


def money(v):
    n = int(round(v))
    return ("-" if n < 0 else "") + "$" + format(abs(n), ",")


def gval(g, raw):
    v = raw.get(g)
    if g == "SeniorCitizen":
        return vl("Yes" if int(v) == 1 else "No")
    if g == "tenure":
        return "%d" % v
    if g in ("MonthlyCharges", "TotalCharges"):
        return "$%.0f" % v
    return vl(v)


def glabel(g, raw):
    return "%s %s: %s" % (FEAT_ICON.get(g, "•"), L["feat"].get(g, g), gval(g, raw))


def cfg_box(page):
    """Business assumptions. Stored in a plain dict so they survive page switches."""
    cfg = st.session_state.setdefault("_cfg", dict(D.DEFAULT_CFG))
    with st.expander(L["as_title"]):
        c = st.columns(4)
        cfg["horizon"] = int(c[0].number_input(L["as_horizon"], 3, 36, int(cfg["horizon"]), key=page + "_horizon"))
        cfg["disc1"] = float(c[1].number_input(L["as_c1"], 0.0, 50.0, float(cfg["disc1"]), step=1.0, key=page + "_disc1"))
        cfg["disc2"] = float(c[2].number_input(L["as_c2"], 0.0, 50.0, float(cfg["disc2"]), step=1.0, key=page + "_disc2"))
        cfg["autopay_credit"] = float(c[3].number_input(L["as_auto"], 0.0, 200.0, float(cfg["autopay_credit"]),
                                                          step=1.0, key=page + "_auto"))
        c = st.columns(4)
        cfg["tech_cost"] = float(c[0].number_input(L["as_tech"], 0.0, 500.0, float(cfg["tech_cost"]), step=1.0, key=page + "_tech"))
        cfg["sec_cost"] = float(c[1].number_input(L["as_sec"], 0.0, 500.0, float(cfg["sec_cost"]), step=1.0, key=page + "_sec"))
        hi = int(c[2].slider(L["as_gate_hi"], 20, 95, int(cfg["gate_hi"]), key=page + "_ghi"))
        lo = int(c[3].slider(L["as_gate_lo"], 5, 90, int(cfg["gate_lo"]), key=page + "_glo"))
        cfg["gate_hi"], cfg["gate_lo"] = hi, min(lo, hi)
        st.caption(L["as_note"])
    return cfg


def gauge_fig(p, ref=None):
    import plotly.graph_objects as go
    col = "#1f9d55" if p < 0.30 else ("#d98e04" if p < 0.60 else "#d64545")
    ind = dict(
        mode="gauge+number" + ("+delta" if ref is not None else ""), value=p * 100,
        number={"suffix": "%", "valueformat": ".0f", "font": {"size": 44}},
        gauge={"axis": {"range": [0, 100], "tickcolor": "#8da0b8"}, "bar": {"color": col},
               "bgcolor": "#0a1322", "borderwidth": 0,
               "steps": [{"range": [0, 30], "color": "#10301d"}, {"range": [30, 60], "color": "#33290a"},
                         {"range": [60, 100], "color": "#331111"}]})
    if ref is not None:
        ind["delta"] = {"reference": ref * 100, "suffix": "%", "valueformat": ".0f",
                        "increasing": {"color": "#d64545"}, "decreasing": {"color": "#1f9d55"}}
    fig = go.Figure(go.Indicator(**ind))
    fig.update_layout(height=240, margin=dict(l=25, r=25, t=25, b=5), paper_bgcolor="rgba(0,0,0,0)",
                      font={"color": "#ffffff"})
    return fig


def story_fig(r, raw):
    import plotly.graph_objects as go
    groups = D.grouped_shap(r["cols"], r["shap"])
    top = groups[:7]
    base = r["base_val"] * 100
    d = [v * 100 for _, v in top]
    other = r["p"] * 100 - base - sum(d)
    names = [L["st_base"]] + [glabel(g, raw) for g, _ in top]
    vals = [base] + d
    meas = ["absolute"] + ["relative"] * len(d)
    if abs(other) >= 0.5:
        names.append(L["st_other"])
        vals.append(other)
        meas.append("relative")
    names.append(L["st_final"])
    vals.append(r["p"] * 100)
    meas.append("total")
    text = ["%.0f%%" % vals[0]] + ["%+.0f" % v for v in vals[1:-1]] + ["%.0f%%" % vals[-1]]
    fig = go.Figure(go.Waterfall(
        orientation="h", y=names, x=vals, measure=meas, text=text, textposition="outside",
        connector={"line": {"color": "#3a4a63"}},
        increasing={"marker": {"color": "#d64545"}}, decreasing={"marker": {"color": "#1f9d55"}},
        totals={"marker": {"color": "#1e90ff"}}))
    fig.update_yaxes(autorange="reversed")
    fig.update_xaxes(title=L["st_pts"], gridcolor="#16233a")
    fig.update_layout(height=430, margin=dict(l=10, r=40, t=10, b=40), paper_bgcolor="rgba(0,0,0,0)",
                      plot_bgcolor="rgba(0,0,0,0)", font={"color": "#dbe6f5"}, showlegend=False)
    return fig


@st.cache_resource(show_spinner=False)
def get_map_data():
    from sklearn.decomposition import PCA
    b = get_bundle()
    ex = get_explainers()
    Xtr = b["X_train"]
    Z = (Xtr.values.astype(float) - ex.mu) / ex.sd
    pca = PCA(n_components=2, random_state=0).fit(Z)
    idx = np.random.RandomState(0).choice(len(Xtr), min(1500, len(Xtr)), replace=False)
    y = b["y_train"].values if "y_train" in b else None
    return pca, idx, pca.transform(Z[idx]), y


def map_fig(r):
    import plotly.graph_objects as go
    b = get_bundle()
    ex = get_explainers()
    pca, idx, xy, y = get_map_data()
    Xtr = b["X_train"]
    z = ((np.asarray(r["x"], dtype=float) - ex.mu) / ex.sd).reshape(1, -1)
    me = pca.transform(z)[0]
    _, nn = ex.nn.kneighbors(z, n_neighbors=10)
    nn = nn[0]
    npts = pca.transform((Xtr.values[nn].astype(float) - ex.mu) / ex.sd)
    hov = ["tenure %d mo, $%.0f/mo" % (t, m) for t, m in zip(Xtr["tenure"].values[idx], Xtr["MonthlyCharges"].values[idx])]
    fig = go.Figure()
    if y is not None:
        yy = y[idx]
        for val, name, col in [(1, L["mp_churn"], "#d64545"), (0, L["mp_stay"], "#1f9d55")]:
            m = yy == val
            fig.add_trace(go.Scatter(x=xy[m, 0], y=xy[m, 1], mode="markers", name=name,
                                     marker=dict(size=5, color=col, opacity=0.5),
                                     text=[h for h, k in zip(hov, m) if k], hoverinfo="text"))
        stat = L["mp_stat"].format(k=int(np.sum(y[nn])), n=len(nn), p=round(r["p"] * 100))
    else:
        pr = b["model"].predict_proba(Xtr.iloc[idx])[:, 1]
        fig.add_trace(go.Scatter(x=xy[:, 0], y=xy[:, 1], mode="markers", name=L["mp_pred"],
                                 marker=dict(size=5, color=pr, colorscale="RdYlGn_r", opacity=0.6,
                                             colorbar=dict(title=L["mp_pred"])), text=hov, hoverinfo="text"))
        q = b["model"].predict_proba(Xtr.iloc[nn])[:, 1].mean()
        stat = L["mp_stat2"].format(n=len(nn), q=round(q * 100), p=round(r["p"] * 100))
    lx, ly = [], []
    for px, py in npts:
        lx += [me[0], px, None]
        ly += [me[1], py, None]
    fig.add_trace(go.Scatter(x=lx, y=ly, mode="lines", line=dict(color="rgba(255,212,59,.55)", width=1),
                             showlegend=False, hoverinfo="skip"))
    fig.add_trace(go.Scatter(x=npts[:, 0], y=npts[:, 1], mode="markers", name=L["mp_neigh"],
                             marker=dict(size=13, color="rgba(0,0,0,0)", line=dict(color="#ffd43b", width=2))))
    fig.add_trace(go.Scatter(x=[me[0]], y=[me[1]], mode="markers", name=L["mp_you"],
                             marker=dict(size=24, symbol="star", color="#22b8ff", line=dict(color="#ffffff", width=2))))
    fig.update_xaxes(visible=False)
    fig.update_yaxes(visible=False)
    fig.update_layout(height=470, margin=dict(l=5, r=5, t=5, b=5), paper_bgcolor="rgba(0,0,0,0)",
                      plot_bgcolor="rgba(11,18,32,.6)", font={"color": "#dbe6f5"},
                      legend=dict(orientation="h", y=-0.02))
    return fig, stat


def speak_buttons(items, lang_code):
    """Browser text-to-speech buttons. items = [(label, text)]. Nothing to install."""
    import json
    import streamlit.components.v1 as components
    texts = json.dumps([t for _, t in items]).replace("</", "<\\/")
    btns = "".join("<button onclick='say(%d)'>🔊 %s</button>" % (i, lab) for i, (lab, _) in enumerate(items))
    html = (
        "<style>body{margin:0;font-family:sans-serif}button{background:#1e90ff;color:#fff;border:0;border-radius:999px;"
        "padding:10px 18px;margin:4px 8px 4px 0;font-weight:700;cursor:pointer}"
        "button.s{background:#31435f}</style>%s<button class='s' onclick='speechSynthesis.cancel()'>⏹ %s</button>"
        "<script>const T=%s;const LG=%s;"
        "function pick(){const vs=speechSynthesis.getVoices();const f=LG.toLowerCase();const p=f.slice(0,2);"
        "return vs.find(x=>x.lang.replace('_','-').toLowerCase()===f)||vs.find(x=>x.lang.toLowerCase().startsWith(p));}"
        "function say(i){speechSynthesis.cancel();const u=new SpeechSynthesisUtterance(T[i]);u.lang=LG;"
        "const v=pick();if(v)u.voice=v;u.rate=0.95;speechSynthesis.speak(u);}"
        "speechSynthesis.getVoices();</script>" % (btns, L["ms_stop"], texts, json.dumps(lang_code)))
    try:
        components.html(html, height=70)
    except Exception as e:  # noqa
        st.warning("Voice: %s" % e)


def set_example(kind):
    for k, v in EX[kind].items():
        st.session_state["in_" + k] = v
    st.session_state.pop("res", None)


def build_raw():
    g = st.session_state
    raw = {k: g["in_" + k] for k in EX["high"]}
    raw["SeniorCitizen"] = 1 if raw["SeniorCitizen"] == "Yes" else 0
    if raw["InternetService"] == "No":
        for c in ["OnlineSecurity", "OnlineBackup", "DeviceProtection", "TechSupport", "StreamingTV", "StreamingMovies"]:
            raw[c] = "No internet service"
    elif any(raw[c] == "No internet service" for c in ["OnlineSecurity", "TechSupport"]):
        for c in ["OnlineSecurity", "OnlineBackup", "DeviceProtection", "TechSupport", "StreamingTV", "StreamingMovies"]:
            if raw[c] == "No internet service":
                raw[c] = "No"
    if raw["PhoneService"] == "No":
        raw["MultipleLines"] = "No phone service"
    raw["MonthlyCharges"] = float(raw["MonthlyCharges"])
    raw["tenure"] = int(raw["tenure"])
    raw["TotalCharges"] = round(raw["tenure"] * raw["MonthlyCharges"], 2)
    return raw


def run_analysis(raw):
    from common import encode
    from engine import top_factors
    b = get_bundle()
    ex = get_explainers()
    X = encode(pd.DataFrame([raw]), b["columns"])
    res = ex.explain(X)
    up, down = top_factors(res, 3)
    p = float(D.predict_batch(b, pd.DataFrame([raw]))[0])
    return {"p": p, "up": up, "down": down, "trust": res["trust"], "overlap": res["overlap"],
            "neigh": res["neigh_sim"], "shap": res["shap"], "cols": res["cols"], "x": res["values"],
            "base_val": res["base"], "thr": b["threshold"]}


def pdf_info(r, raw, rows, best, dec, rar, cfg, fair):
    """Everything the PDF needs, in English."""
    E = {**T["en"], **XT["en"]}
    ev = lambda v: str(v)  # noqa: E731
    lvl = risk_level_fn(r["p"])
    groups = D.grouped_shap(r["cols"], r["shap"])
    top = groups[:6]
    story = []
    for g, v in top:
        val = ("Yes" if int(raw[g]) == 1 else "No") if g == "SeniorCitizen" else (
            "%d" % raw[g] if g == "tenure" else ("$%.0f" % raw[g] if g in ("MonthlyCharges", "TotalCharges") else ev(raw[g])))
        story.append(("%s: %s" % (E["feat"].get(g, g), val), v * 100))
    other = r["p"] * 100 - r["base_val"] * 100 - sum(v for _, v in story)
    if abs(other) >= 0.5:
        story.append((E["st_other"], other))
    tl = "high" if r["trust"] >= 70 else ("medium" if r["trust"] >= 40 else "low")
    offer = E["scen"][best["key"]] if (best and dec in ("act", "review")) else None
    if offer:
        sms = E["sms_offer"].format(name="Customer", offer=offer)
        if dec == "review":
            sms = E["ms_draft"] + " " + sms
    else:
        sms = E["sms_check"].format(name="Customer")
    prof = ("Tenure %d months | Monthly charges $%.0f | Contract: %s | Internet: %s | Payment: %s | Senior citizen: %s | "
            "Tech support: %s | Online security: %s") % (
        raw["tenure"], raw["MonthlyCharges"], raw["Contract"], raw["InternetService"], raw["PaymentMethod"],
        "Yes" if raw["SeniorCitizen"] == 1 else "No", raw["TechSupport"], raw["OnlineSecurity"])
    return {
        "p": r["p"], "level": E["risk"][lvl], "trust": r["trust"], "trust_level": E["tr_level"][tl],
        "horizon": cfg["horizon"], "rar": rar, "dec": dec, "dec_name": E["dec_names"][dec], "dec_body": E["dec_body"][dec],
        "profile": prof, "base_risk": r["base_val"] * 100, "story": story,
        "story_base_label": E["st_base"], "story_final_label": E["st_final"],
        "rows": [{"label": E["scen"][x["key"]], **{k: x[k] for k in ("new_p", "delta", "cost", "saved", "net")}}
                 for x in rows if x["delta"] > 0.001],
        "trust_line": "Explanation trust %d/100 (%s): SHAP vs LIME top-5 agreement %d%%, similarity with nearest customers %d%%." % (
            r["trust"], E["tr_level"][tl], round(r["overlap"] * 100), round(r["neigh"] * 100)),
        "fair_line": "Fairness: flipping the senior-citizen flag changes this prediction by %+.1f pts; gender is not a model input." % (
            fair[2] * 100),
        "sms": sms,
        "footnote": "Risk alert, not a verdict: a human makes the final decision. Offer effects come from a model what-if, not proof "
                    "of cause; pilot before rollout. Amounts are in dataset currency units and use the business assumptions set in the app.",
    }


def risk_level_fn(p):
    from engine import risk_level
    return risk_level(p)


# ------------------------------------------------------------------ ANALYZE
EX = {
    "high": dict(SeniorCitizen="No", Partner="No", Dependents="No", tenure=3, PhoneService="Yes",
                 MultipleLines="No", InternetService="Fiber optic", OnlineSecurity="No", OnlineBackup="No",
                 DeviceProtection="No", TechSupport="No", StreamingTV="No", StreamingMovies="No",
                 Contract="Month-to-month", PaperlessBilling="Yes", PaymentMethod="Electronic check",
                 MonthlyCharges=85.0),
    "low": dict(SeniorCitizen="No", Partner="Yes", Dependents="Yes", tenure=60, PhoneService="Yes",
                MultipleLines="No", InternetService="DSL", OnlineSecurity="Yes", OnlineBackup="Yes",
                DeviceProtection="No", TechSupport="Yes", StreamingTV="No", StreamingMovies="No",
                Contract="Two year", PaperlessBilling="No", PaymentMethod="Bank transfer (automatic)",
                MonthlyCharges=55.0),
}


def page_analyze():
    b = get_bundle()
    opts = b["options"]
    for _k, _v in EX["high"].items():      # fill any missing form field with the high-risk example
        st.session_state.setdefault("in_" + _k, _v)
    st.subheader(L["an_title"])
    st.caption(L["an_sub"])
    e1, e2, _ = st.columns([1.3, 1.3, 3])
    e1.button(L["ex_high"], on_click=set_example, args=("high",), width="stretch")
    e2.button(L["ex_low"], on_click=set_example, args=("low",), width="stretch")

    F = L["fields"]
    yn = ["Yes", "No"]
    c1, c2, c3 = st.columns(3)
    with c1:
        st.selectbox(F["SeniorCitizen"], ["No", "Yes"], key="in_SeniorCitizen", format_func=vl)
        st.selectbox(F["Partner"], yn, key="in_Partner", format_func=vl)
        st.selectbox(F["Dependents"], yn, key="in_Dependents", format_func=vl)
        st.slider(F["tenure"], 0, b["tenure_max"], key="in_tenure")
        st.number_input(F["MonthlyCharges"], 10.0, 200.0, step=1.0, key="in_MonthlyCharges")
        st.selectbox(F["PhoneService"], yn, key="in_PhoneService", format_func=vl)
    with c2:
        for k in ["InternetService", "MultipleLines", "OnlineSecurity", "OnlineBackup", "DeviceProtection"]:
            st.selectbox(F[k], opts[k], key="in_" + k, format_func=vl)
    with c3:
        for k in ["TechSupport", "StreamingTV", "StreamingMovies", "Contract", "PaymentMethod"]:
            st.selectbox(F[k], opts[k], key="in_" + k, format_func=vl)
        st.selectbox(F["PaperlessBilling"], yn, key="in_PaperlessBilling", format_func=vl)

    cfg = cfg_box("an")

    if st.button(L["run"], type="primary"):
        with st.spinner(L["spin"]):
            raw = build_raw()
            st.session_state.res = run_analysis(raw)
            st.session_state.res_raw = raw
            st.session_state.aid = st.session_state.get("aid", 0) + 1

    r = st.session_state.get("res")
    if not r:
        st.info(L["need_run"])
        return
    raw = st.session_state.res_raw
    aid = st.session_state.get("aid", 0)
    st.divider()

    # ---- money + decision (recomputed live so edited assumptions apply immediately)
    base_p, plan = D.plan_actions(b, pd.DataFrame([raw]), cfg)
    p = r["p"]
    rows = D.single_rows(plan)
    best = D.best_of(rows)
    rar = float(D.revenue_at_risk(p, raw["MonthlyCharges"], cfg["horizon"]))
    dec = D.decide(p, r["trust"], best["net"] if best else None, cfg)
    lvl = risk_level_fn(p)

    st.markdown("<div class='dec dec-%s'><div class='dec-t'>%s</div><div class='dec-n'>%s</div>"
                "<div class='dec-b'>%s</div></div>" % (dec, L["dec_title"], L["dec_names"][dec], L["dec_body"][dec]),
                unsafe_allow_html=True)

    tabs = st.tabs(L["tabs"])

    # ============ 1. DECISION
    with tabs[0]:
        g1, g2, g3, g4 = st.columns([1.25, 1, 1.2, 1])
        with g1:
            st.plotly_chart(gauge_fig(p), width="stretch", key="gauge_main_%d" % aid)
        kpi = "<div class='kpi'><small>%s</small><b>%s</b></div>"
        g2.markdown(kpi % (L["prob"], "%d%%" % round(p * 100)) +
                    "<span class='pill %s'>%s</span>" % (lvl, L["risk"][lvl]), unsafe_allow_html=True)
        g2.caption("%s: %d%%" % (L["thr_note"], round(r["thr"] * 100)))
        g3.markdown(kpi % (L["m_rar"], money(rar)), unsafe_allow_html=True)
        g3.caption(L["m_rar_help"])
        g4.markdown(kpi % (L["m_net"], money(best["net"]) if best else "-"), unsafe_allow_html=True)
        g4.caption(L["scen"][best["key"]] if best else L["m_none"])

        up = ", ".join(flabel(f) for f, _ in r["up"]) or L["sum_none"]
        dn = ", ".join(flabel(f) for f, _ in r["down"])
        text = L["sum_tpl"].format(p=round(p * 100), risk=L["risk"][lvl], up=up,
                                   down=L["sum_down"].format(down=dn) if dn else "")
        st.markdown("**%s**" % L["sum_title"])
        st.info(text)

        st.subheader(L["wi_title"])
        st.caption(L["wi_intro"])
        good = [x for x in rows if x["delta"] >= 0.005]
        if not good:
            st.warning(L["wi_none"])
        else:
            cols = L["wi_cols2"]
            tab = [{cols[0]: L["scen"][x["key"]], cols[1]: "%d%%" % round(x["new_p"] * 100),
                    cols[2]: "-%d pts" % round(x["delta"] * 100), cols[3]: money(x["cost"]),
                    cols[4]: money(x["saved"]), cols[5]: money(x["net"])} for x in good]
            st.dataframe(pd.DataFrame(tab), hide_index=True, width="stretch")
        st.caption(L["wi_caveat"])

        st.subheader(L["tr_title"])
        t = r["trust"]
        tl = "high" if t >= 70 else ("medium" if t >= 40 else "low")
        t1, t2, t3 = st.columns(3)
        t1.metric(L["tr_title"], "%d / 100" % t, L["tr_level"][tl], delta_color="off")
        t2.metric(L["tr_a"], "%d%%" % round(r["overlap"] * 100))
        t3.metric(L["tr_b"], "%d%%" % round(r["neigh"] * 100))
        st.progress(t / 100)
        st.caption(L["tr_help"])

        fair = D.senior_counterfactual(b, raw)
        try:
            from report import build_pdf
            pk = (aid, tuple(sorted(cfg.items())), dec)
            if st.session_state.get("_pdfk") != pk:   # build once, not on every live-slider rerun
                st.session_state["_pdf"] = build_pdf(pdf_info(r, raw, rows, best, dec, rar, cfg, fair))
                st.session_state["_pdfk"] = pk
            pdf = st.session_state["_pdf"]
            st.download_button("📄 " + L["pdf_btn"], pdf, file_name="churnlens_customer_report.pdf",
                               mime="application/pdf", key="pdf_%d" % aid)
            st.caption(L["pdf_note"])
        except Exception as e:  # noqa
            st.warning("PDF: %s" % e)
        st.markdown("<div class='note'>%s</div>" % L["note"], unsafe_allow_html=True)

    # ============ 2. STORY (waterfall)
    with tabs[1]:
        st.subheader(L["st_title"])
        st.caption(L["st_sub"])
        st.plotly_chart(story_fig(r, raw), width="stretch", key="story_%d" % aid)

    # ============ 3. LIVE WHAT-IF
    with tabs[2]:
        st.subheader(L["lv_title"])
        st.caption(L["lv_sub"])
        internet = raw["InternetService"] != "No"
        new = dict(raw)
        w1, w2 = st.columns(2)
        with w1:
            co = opts["Contract"]
            new["Contract"] = st.radio(L["lv_contract"], co, index=co.index(raw["Contract"]), format_func=vl,
                                       horizontal=True, key="lv_c_%d" % aid)
            pm = opts["PaymentMethod"]
            new["PaymentMethod"] = st.selectbox(L["lv_pay"], pm, index=pm.index(raw["PaymentMethod"]),
                                                format_func=vl, key="lv_p_%d" % aid)
        with w2:
            if internet:
                new["TechSupport"] = "Yes" if st.toggle(L["lv_tech"], value=raw["TechSupport"] == "Yes",
                                                        key="lv_t_%d" % aid) else "No"
                new["OnlineSecurity"] = "Yes" if st.toggle(L["lv_sec"], value=raw["OnlineSecurity"] == "Yes",
                                                           key="lv_s_%d" % aid) else "No"
            else:
                st.info(L["lv_nointernet"])
        p2 = float(D.predict_batch(b, pd.DataFrame([new]))[0])
        cost = D.scenario_cost(raw, new, cfg)
        saved = (p - p2) * raw["MonthlyCharges"] * cfg["horizon"]
        l1, l2 = st.columns([1.3, 1])
        with l1:
            st.plotly_chart(gauge_fig(p2, ref=p), width="stretch", key="gauge_live_%d" % aid)
        with l2:
            kpi = "<div class='kpi'><small>%s</small><b>%s</b></div>"
            st.markdown(kpi % (L["lv_orig"], "%d%%" % round(p * 100)) + kpi % (L["lv_new"], "%d%%" % round(p2 * 100)),
                        unsafe_allow_html=True)
        k1, k2, k3 = st.columns(3)
        k1.metric(L["lv_saved"], money(saved))
        k2.metric(L["lv_cost"], money(cost))
        k3.metric(L["lv_net"], money(saved - cost))
        st.caption(L["wi_caveat"])

    # ============ 4. SIMILAR CUSTOMERS MAP
    with tabs[3]:
        st.subheader(L["mp_title"])
        st.caption(L["mp_sub"])
        try:
            fig, stat = map_fig(r)
            st.info(stat)
            st.plotly_chart(fig, width="stretch", key="map_%d" % aid)
        except Exception as e:  # noqa
            st.warning("Map: %s" % e)

    # ============ 5. MESSAGE + VOICE
    with tabs[4]:
        st.subheader(L["ms_title"])
        m1, m2 = st.columns([1, 2])
        nm = m1.text_input(L["ms_name"], value=DEFAULT_NAME[lang], key="ms_nm")
        ml = m2.segmented_control(L["ms_lang"], list(LANG_NAMES), default=lang, key="ms_lang_sel",
                                  format_func=lambda c: LANG_NAMES[c]) or lang
        X2, L2 = XT[ml], T[ml]
        offer = L2["scen"][best["key"]] if (best and dec in ("act", "review")) else None
        if offer:
            sms = X2["sms_offer"].format(name=nm, offer=offer)
            call = X2["call_offer"].format(name=nm, offer=offer)
            if dec == "review":
                sms, call = X2["ms_draft"] + " " + sms, X2["ms_draft"] + " " + call
            extra = X2["sp_extra_offer"].format(offer=offer, net=money(best["net"]))
        else:
            sms = X2["sms_check"].format(name=nm)
            call = X2["call_check"].format(name=nm)
            extra = X2["sp_extra_none"]
        spoken = X2["sp_sum"].format(name=nm, p=round(p * 100), dec=X2["dec_names"][dec], extra=extra)
        st.markdown("**%s**" % L["ms_sms"])
        st.code(sms, language=None)
        st.markdown("**%s**" % L["ms_call"])
        st.code(call, language=None)
        st.markdown("**🔊 %s**" % L["ms_listen"])
        speak_buttons([(L["ms_listen_sum"], spoken), (L["ms_listen_call"], call)], VOICE_LANG[ml])
        st.caption(L["ms_tts_note"])

    # ============ 6. FAIRNESS (customer level)
    with tabs[5]:
        st.subheader(L["fr_title"])
        st.caption(L["fr_sub"])
        p0, p1, d = D.senior_counterfactual(b, raw)
        f1, f2, f3, f4 = st.columns(4)
        f1.metric(L["fr_orig"], "%.1f%%" % (p0 * 100))
        f2.metric("%s (%s)" % (L["fr_flip"], L["feat"]["SeniorCitizen"]), "%.1f%%" % (p1 * 100))
        f3.metric(L["fr_delta"], "%+.1f pts" % (d * 100), delta_color="off")
        f4.metric(L["fr_gender"], L["fr_gender_val"])
        if abs(d) * 100 >= 5:
            st.warning(L["fr_warn"].format(d="%+.1f" % (d * 100)))
        else:
            st.success(L["fr_ok"])
        st.caption(L["fr_note"])


# ------------------------------------------------------------------ BATCH
def make_xlsx(show, summary):
    from openpyxl.styles import Font, PatternFill
    buf = io.BytesIO()
    with pd.ExcelWriter(buf, engine="openpyxl") as w:
        summary.to_excel(w, sheet_name=L["bt_xl_sum"][:31], index=False)
        show.to_excel(w, sheet_name="Customers", index=False)
        for ws in w.book.worksheets:
            for c in ws[1]:
                c.font = Font(bold=True, color="FFFFFF")
                c.fill = PatternFill("solid", fgColor="0B3D75")
            for col in ws.columns:
                width = max(len(str(c.value)) if c.value is not None else 0 for c in list(col)[:80]) + 2
                ws.column_dimensions[col[0].column_letter].width = min(42, width)
            ws.freeze_panes = "A2"
    return buf.getvalue()


@st.cache_resource(show_spinner=False)
def get_batch_expl():
    return D.BatchExplainer(get_bundle())


def page_batch():
    from common import encode
    b = get_bundle()
    st.subheader(L["bt_title"])
    st.caption(L["bt_sub"])
    cfg = cfg_box("bt")

    sample = D.sample_csv_bytes()
    if sample:
        st.download_button("⬇ " + L["bt_sample"], sample, file_name="sample_customers.csv", mime="text/csv")
    up = st.file_uploader(L["bt_upload"], type=["csv"])
    if up is None:
        st.info(L["bt_info"])
        return
    try:
        df, missing = D.prepare_batch(pd.read_csv(up))
    except Exception as e:  # noqa
        st.error(str(e))
        return
    if missing:
        st.error(L["bt_missing"].format(cols=", ".join(missing)))
        return
    st.caption(L["bt_rows"].format(n=len(df)))
    full_n = st.slider(L["bt_fulln"], 0, 30, 0, key="bt_fulln")

    if st.button(L["bt_run"], type="primary"):
        with st.spinner(L["bt_spin"]):
            res = D.run_batch(b, get_batch_expl(), df, cfg)
        if full_n > 0:
            ex = get_explainers()
            top = res[res["decision"] != "monitor"].sort_values("rar", ascending=False).head(full_n).index
            bar = st.progress(0.0)
            for t, i in enumerate(top):
                bar.progress((t + 1) / max(len(top), 1), text=L["bt_full_spin"].format(i=t + 1, n=len(top)))
                er = ex.explain(encode(df.iloc[[i]], b["columns"]))
                res.loc[i, "trust"] = er["trust"]
                res.loc[i, "trust_type"] = "full"
            bar.empty()
            res = D.apply_decisions(res, cfg)
        st.session_state.batch = res
        st.session_state.batch_cfg = dict(cfg)

    res = st.session_state.get("batch")
    if res is None:
        return
    bc = st.session_state.get("batch_cfg", cfg)
    st.divider()

    high = int((res["p"] >= 0.60).sum())
    gain = float(res.loc[res["decision"] == "act", "net"].sum())
    k1, k2, k3, k4 = st.columns(4)
    k1.metric(L["bt_n"], format(len(res), ","))
    k2.metric(L["bt_high"], format(high, ","))
    k3.metric(L["bt_rev"], money(res["rar"].sum()))
    k4.metric(L["bt_gain"], money(gain))

    counts = res["decision"].value_counts()
    chart = pd.DataFrame({"n": [int(counts.get(d, 0)) for d in DEC_ORDER]},
                         index=[L["dec_names"][d] for d in DEC_ORDER])
    st.markdown("**%s**" % L["bt_dec_chart"])
    st.bar_chart(chart, horizontal=True)

    so = L["bt_sort_opts"]
    key = st.selectbox(L["bt_sort"], list(so), format_func=lambda k: so[k], key="bt_sortby")
    o = res.sort_values(key, ascending=False).reset_index(drop=True)
    C = L["bt_cols"]
    show = pd.DataFrame({
        C["rank"]: np.arange(1, len(o) + 1),
        C["id"]: o["id"].values,
        C["p"]: (o["p"] * 100).round(0).values,
        C["level"]: [L["risk"][D.risk_word(p, bc)] for p in o["p"]],
        C["monthly"]: o["monthly"].round(0).values,
        C["rar"]: o["rar"].round(0).values,
        C["reason1"]: [flabel(c) if isinstance(c, str) else "-" for c in o["reason1"]],
        C["reason2"]: [flabel(c) if isinstance(c, str) else "-" for c in o["reason2"]],
        C["best"]: [L["scen"][k] if isinstance(k, str) else "-" for k in o["best"]],
        C["cost"]: o["cost"].round(0).values,
        C["saved"]: o["saved"].round(0).values,
        C["net"]: o["net"].round(0).values,
        C["trust"]: o["trust"].values,
        C["trust_type"]: [L["bt_tt"][t] for t in o["trust_type"]],
        C["decision"]: [L["dec_names"][d] for d in o["decision"]],
    })
    st.dataframe(show, hide_index=True, width="stretch", height=460, column_config={
        C["p"]: st.column_config.ProgressColumn(C["p"], min_value=0, max_value=100, format="%d"),
        C["trust"]: st.column_config.ProgressColumn(C["trust"], min_value=0, max_value=100, format="%d"),
    })
    st.caption(L["bt_note"])

    summary = pd.DataFrame(
        [[L["bt_n"], len(res)], [L["bt_high"], high], [L["bt_rev"], round(float(res["rar"].sum()))],
         [L["bt_gain"], round(gain)]] + [[L["dec_names"][d], int(counts.get(d, 0))] for d in DEC_ORDER],
        columns=["", ""])
    d1, d2, _ = st.columns([1, 1.3, 2])
    d1.download_button("⬇ " + L["bt_dl_csv"], show.to_csv(index=False).encode("utf-8-sig"),
                       file_name="churnlens_batch_report.csv", mime="text/csv")
    try:
        d2.download_button("⬇ " + L["bt_dl_xlsx"], make_xlsx(show, summary), file_name="churnlens_batch_report.xlsx",
                           mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
    except Exception as e:  # noqa
        d2.warning("Excel: %s" % e)


# ------------------------------------------------------------------ MODEL & FAIRNESS
def page_model():
    b = get_bundle()
    st.subheader(L["md_title"])
    st.write("**%s:** %s" % (L["md_chosen"], b["model_name"]))
    st.dataframe(b["metrics"], width="stretch")
    st.caption(L["md_final"])
    st.json(b["final_metrics"])

    st.subheader(L["gi_title"])
    if st.button(L["gi_btn"]):
        with st.spinner(L["spin"]):
            gi = get_explainers().global_importance().head(12)
            st.session_state.gi = gi
    gi = st.session_state.get("gi")
    if gi is not None:
        g = pd.DataFrame({"mean |SHAP|": gi.values}, index=[flabel(i) for i in gi.index])
        st.bar_chart(g, horizontal=True)

    st.subheader(L["fa_title"])
    st.caption(L["fa_note"])
    fa = b["fairness"]
    st.dataframe(fa, hide_index=True, width="stretch")
    for gt, sub in fa.groupby("Group type"):
        gap = sub["Recall"].max() - sub["Recall"].min()
        st.write("%s (%s): **%.1f pts**" % (L["fa_gap"], gt, gap * 100))


# ------------------------------------------------------------------ ABOUT
def page_about():
    c1, c2 = st.columns([1, 2])
    with c1:
        if os.path.exists(LOGO):
            st.image(LOGO, width="stretch")
    with c2:
        st.subheader(L["ab_title"])
        st.write(L["ab_body"])
        st.markdown("**%s**" % L["ab_limits_t"])
        st.write(L["ab_limits"])
        st.markdown("`Python · pandas · scikit-learn · XGBoost · SHAP · LIME · Streamlit`")


try:
    {"home": page_home, "analyze": page_analyze, "batch": page_batch, "model": page_model, "about": page_about}[PAGE]()
except FileNotFoundError:
    st.error("models/bundle.joblib not found. Run:  python train.py   (then restart the app)")

st.markdown("<div class='foot'>%s</div>" % L["footer"], unsafe_allow_html=True)
